import os
import io
import logging
from typing import List, Dict, Any

logger = logging.getLogger("sovereign.rag.loaders")

class TextLoader:
    """Loader for plain text and markdown documents."""
    def load(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found: {file_path}")

        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1", errors="replace") as f:
                text = f.read()

        return [{"text": text, "metadata": {"source": file_path, "page": 1, "format": "text"}}]


class PDFLoader:
    """
    Production loader for PDF documents using PyPDF2 with page-by-page metadata
    and local OCR fallback for scanned pages.
    """
    def load(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found: {file_path}")

        import PyPDF2
        pages_content: List[Dict[str, Any]] = []

        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                total_pages = len(reader.pages)

                for idx, page in enumerate(reader.pages):
                    page_num = idx + 1
                    raw_text = page.extract_text() or ""
                    clean_text = raw_text.strip()
                    ocr_applied = False

                    # If page has very little selectable text, attempt OCR if images exist
                    if len(clean_text) < 30:
                        ocr_text = self._try_ocr_page(page)
                        if ocr_text:
                            clean_text = ocr_text
                            ocr_applied = True

                    if clean_text:
                        pages_content.append({
                            "text": clean_text,
                            "metadata": {
                                "source": file_path,
                                "page": page_num,
                                "total_pages": total_pages,
                                "ocr_applied": ocr_applied,
                                "format": "pdf"
                            }
                        })
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}")
            raise

        return pages_content

    def _try_ocr_page(self, page) -> str:
        """Attempt OCR on embedded images in the PDF page."""
        try:
            from PIL import Image
            import pytesseract

            extracted_chunks = []
            if hasattr(page, "images") and page.images:
                for img_file in page.images:
                    img_bytes = img_file.data
                    img = Image.open(io.BytesIO(img_bytes))
                    txt = pytesseract.image_to_string(img)
                    if txt.strip():
                        extracted_chunks.append(txt.strip())
            return "\n".join(extracted_chunks)
        except Exception as e:
            logger.debug(f"OCR fallback skipped or failed: {e}")
            return ""


class DOCXLoader:
    """Loader for Word documents with paragraph and table extraction."""
    def load(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found: {file_path}")

        import docx
        doc = docx.Document(file_path)
        content_parts = []

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                content_parts.append(p.text.strip())

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                content_parts.append("\n".join(table_rows))

        full_text = "\n\n".join(content_parts)
        return [{
            "text": full_text,
            "metadata": {"source": file_path, "page": 1, "format": "docx"}
        }]


class ImageOCRLoader:
    """Loader for standalone images (P&ID diagrams, scanned docs) using OCR."""
    def load(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found: {file_path}")

        from PIL import Image
        import pytesseract

        try:
            img = Image.open(file_path)
            ocr_text = pytesseract.image_to_string(img)
            return [{
                "text": ocr_text.strip(),
                "metadata": {
                    "source": file_path,
                    "page": 1,
                    "ocr_applied": True,
                    "format": "image"
                }
            }]
        except Exception as e:
            logger.error(f"Image OCR failed for {file_path}: {e}")
            return [{
                "text": f"[OCR failed for image: {os.path.basename(file_path)}]",
                "metadata": {"source": file_path, "page": 1, "ocr_applied": False}
            }]


def load_document(file_path: str) -> List[Dict[str, Any]]:
    """Universal document loader dispatching by file extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return PDFLoader().load(file_path)
    elif ext in [".docx", ".doc"]:
        return DOCXLoader().load(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
        return ImageOCRLoader().load(file_path)
    else:
        return TextLoader().load(file_path)

